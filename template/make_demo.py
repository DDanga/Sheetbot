"""Builds sheetbot-demo.xlsx: a made-up corner shop, July to October 2026.

    python template/make_demo.py

The numbers are random but fixed (same seed, same file), and shaped so every
feature has something to show:
- Stock: 4 of 15 products are below their minimum;
- Sales: a summer slump in August, growth in September, a weak last week of
  September (for "sales fell by more than 20%"), and a strong October;
- coffee beans sell rarely but earn the most: the "hidden star".
"""

from __future__ import annotations

import random
from datetime import date, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

OUT = Path(__file__).with_name("sheetbot-demo.xlsx")
FIRST, LAST = date(2026, 7, 1), date(2026, 10, 31)

# product, category, in stock, minimum, price, purchase price, supplier, how often it sells, usual quantity
PRODUCTS = [
    ("Whole milk", "Dairy", 3, 10, 3.29, 2.40, "Green Valley Farm", 10, (1, 4)),
    ("Greek yogurt", "Dairy", 14, 8, 1.99, 1.20, "Green Valley Farm", 9, (2, 6)),
    ("Gouda cheese", "Dairy", 6, 5, 8.90, 5.40, "Cheese Corner", 3, (1, 2)),
    ("Sourdough bread", "Bakery", 2, 12, 5.50, 3.20, "Morning Bakery", 7, (1, 3)),
    ("Baguette", "Bakery", 9, 6, 2.80, 1.60, "Morning Bakery", 9, (1, 4)),
    ("Croissant", "Bakery", 18, 10, 2.20, 0.80, "Morning Bakery", 6, (2, 6)),
    ("Coffee beans", "Pantry", 7, 4, 14.90, 6.80, "Bean & Leaf Roasters", 4, (1, 2)),
    ("Green tea", "Pantry", 11, 5, 6.50, 2.60, "Bean & Leaf Roasters", 3, (1, 3)),
    ("Pasta", "Pantry", 25, 10, 1.80, 1.10, "Agro Plus", 6, (1, 4)),
    ("Basmati rice", "Pantry", 4, 8, 3.40, 2.30, "Agro Plus", 4, (1, 3)),
    ("Apples", "Fruit", 40, 15, 2.99, 2.00, "Orchard Fresh", 7, (1, 4)),
    ("Bananas", "Fruit", 22, 15, 1.49, 1.05, "Tropic Imports", 8, (2, 6)),
    ("Oranges", "Fruit", 5, 12, 3.49, 2.40, "Tropic Imports", 5, (1, 4)),
    ("Sparkling water", "Drinks", 60, 24, 0.99, 0.55, "Spring Water Co.", 8, (2, 8)),
    ("Orange juice", "Drinks", 13, 10, 3.90, 2.70, "Spring Water Co.", 5, (1, 3)),
]
MANAGERS = (("Anna", 5), ("Ben", 3), ("Chloe", 2))

CLIENTS = [
    ("Olivia Brooks", "Portland", date(2026, 10, 30), 1240), ("Ethan Clarke", "Seattle", date(2026, 10, 18), 830),
    ("Maya Patel", "Portland", date(2026, 8, 14), 210), ("Lucas Moreau", "Eugene", date(2026, 10, 27), 1560),
    ("Sofia Reyes", "Salem", date(2026, 7, 9), 95), ("Noah Fischer", "Portland", date(2026, 9, 25), 670),
    ("Emma Larsen", "Seattle", date(2026, 10, 5), 2310), ("Liam O'Connor", "Bend", date(2026, 9, 2), 140),
    ("Hana Sato", "Portland", date(2026, 10, 22), 980), ("Daniel Kim", "Eugene", date(2026, 8, 30), 455),
]

MONEY = '"$"#,##0.00'
DATE = "yyyy-mm-dd"
HEADER_FILL = PatternFill("solid", start_color="DCE6F1")


def sales_per_day(day: date, rng: random.Random) -> int:
    """How busy the shop is: weekends, the August slump, a bad week, a good October."""
    base = 9 if day.weekday() >= 5 else 6
    season = {7: 1.0, 8: 0.85, 9: 1.1, 10: 1.3}[day.month]
    if date(2026, 9, 24) <= day <= date(2026, 9, 30):
        season = 0.7
    return max(1, round(base * season * rng.uniform(0.8, 1.2)))


def sales_rows(rng: random.Random) -> list[list]:
    weights = [p[7] for p in PRODUCTS]
    rows = []
    day = FIRST
    while day <= LAST:
        for _ in range(sales_per_day(day, rng)):
            name, _, _, _, price, purchase, _, _, (lo, hi) = rng.choices(PRODUCTS, weights)[0]
            qty = rng.randint(lo, hi)
            manager = rng.choices([m for m, _ in MANAGERS], [w for _, w in MANAGERS])[0]
            rows.append([day, name, qty, round(qty * price, 2), round(qty * purchase, 2), manager])
        day += timedelta(days=1)
    return rows


def sheet(wb: Workbook, title: str, headers: list[str], rows: list[list], widths: list[int],
          formats: dict[int, str]) -> None:
    ws = wb.create_sheet(title)
    ws.append(headers)
    for row in rows:
        ws.append(row)
    for cell in ws[1]:
        cell.font = Font(name="Arial", bold=True)
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(vertical="center")
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="Arial")
            if cell.column in formats:
                cell.number_format = formats[cell.column]
    for i, width in enumerate(widths, start=1):
        ws.column_dimensions[chr(64 + i)].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def main() -> None:
    rng = random.Random(2026)
    wb = Workbook()
    wb.remove(wb.active)
    sheet(wb, "Stock",
          ["Product", "Category", "In stock", "Minimum", "Price", "Purchase price", "Supplier"],
          [list(p[:7]) for p in PRODUCTS], [20, 12, 10, 10, 10, 15, 22], {5: MONEY, 6: MONEY})
    sheet(wb, "Sales",
          ["Date", "Product", "Quantity", "Amount", "Cost", "Manager"],
          sales_rows(rng), [12, 20, 10, 12, 12, 10], {1: DATE, 4: MONEY, 5: MONEY})
    sheet(wb, "Clients",
          ["Client", "City", "Last purchase", "Total spent"],
          [list(c) for c in CLIENTS], [18, 12, 15, 13], {3: DATE, 4: MONEY})
    wb.properties.creator = "Sheetbot demo"
    wb.save(OUT)
    print(f"Saved {OUT}")


if __name__ == "__main__":
    main()
