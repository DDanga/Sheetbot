"""A tiny fake Telegram chat for driving the bot like a user would."""
import asyncio
from pathlib import Path
from types import SimpleNamespace
from zoneinfo import ZoneInfo

import openpyxl

from sheetbot.config import Config
from sheetbot.sheets import table_from
from sheetbot.storage import Storage

HERE = Path(__file__).resolve().parent
# the original Russian demo shop (September 2026): most tests are pinned to its numbers
DEMO_XLSX = HERE / "fixtures" / "demo-ru.xlsx"
# the demo users get: an English shop, July–October 2026
TEMPLATE_XLSX = HERE.parent / "template" / "sheetbot-demo.xlsx"


def _shown(cell) -> str:
    """A cell as Google shows it: '$3.29', '2026-09-01' or '01.09.2026', '12'."""
    v, fmt = cell.value, cell.number_format or ""
    if v is None:
        return ""
    if hasattr(v, "strftime"):
        return v.strftime("%Y-%m-%d" if fmt.lower().startswith("yyyy") else "%d.%m.%Y")
    if "$" in fmt and isinstance(v, (int, float)):
        return f"${v:,.2f}"
    return str(v)


def demo_grids(path: Path = DEMO_XLSX) -> dict[str, list[list[str]]]:
    """A demo spreadsheet as Google returns it: every cell as displayed text."""
    wb = openpyxl.load_workbook(path)
    return {ws.title: [[_shown(c) for c in row] for row in ws.iter_rows()] for ws in wb.worksheets}


def template_grids() -> dict[str, list[list[str]]]:
    return demo_grids(TEMPLATE_XLSX)


class FakeSheets:
    def __init__(self, grids=None, error=None, decimal=None):
        self.grids = grids if grids is not None else demo_grids()
        self.error = error
        self.decimal = decimal      # what Google's locale would tell the real client

    def _check(self):
        if self.error:
            raise self.error

    def read_all(self, max_age=0):
        from sheetbot.sheets import summarize
        self._check()
        return summarize("Магазин — демо", self.grids)

    def read_grids(self, max_age=0):
        import copy
        self._check()
        return "Магазин — демо", copy.deepcopy(self.grids)

    def tables(self, max_age=0):
        self._check()
        return {name: table_from(name, v, self.decimal) for name, v in self.grids.items()}

    def table(self, name, max_age=0):
        return self.tables().get(name)


class FakeBot:
    """Collects everything the background check sends, by chat."""

    username = "sheetbot_demo_bot"

    def __init__(self, blocked=()):
        self.sent = []          # [(chat_id, text, markup)]
        self.blocked = set(blocked)

    async def send_message(self, chat_id, text, reply_markup=None):
        from telegram.error import Forbidden
        if chat_id in self.blocked:
            raise Forbidden("bot was blocked by the user")
        self.sent.append((chat_id, text, reply_markup))

    def to(self, chat_id):
        return [text for cid, text, _ in self.sent if cid == chat_id]


class Chat:
    """Records what the bot sends and lets the test press buttons and type."""

    def __init__(self, bot_handlers, context, chat_id=100, name="Danga"):
        self.id = chat_id
        self.name = name
        self.context = context
        self.handlers = bot_handlers
        self.messages = []          # [(text, markup)]
        self.alerts = []

    # --- what the bot calls ---
    async def send_message(self, text, reply_markup=None):
        self.messages.append([text, reply_markup])
        return SimpleNamespace(message_id=len(self.messages))

    # --- what the test calls ---
    @property
    def last(self):
        return self.messages[-1][0]

    @property
    def buttons(self):
        markup = self.messages[-1][1]
        return [b for row in markup.inline_keyboard for b in row] if markup else []

    def button_texts(self):
        return [b.text for b in self.buttons]

    def _update(self, text=None, query=None):
        return SimpleNamespace(
            effective_chat=self, callback_query=query,
            effective_user=SimpleNamespace(full_name=self.name),
            effective_message=SimpleNamespace(text=text),
        )

    def command(self, handler, *args):
        self.context.args = list(args)
        asyncio.run(handler(self._update(), self.context))

    def receive(self, text, markup=None):
        """Put a message the bot sent on its own (via FakeBot) on screen, to press its buttons."""
        self.messages.append([text, markup])

    def another(self, chat_id, name):
        """A second person in the same bot: shared bot and data, their own screen and drafts."""
        context = SimpleNamespace(application=self.context.application, user_data={},
                                  bot=self.context.bot, job_queue=None, args=[])
        return Chat(self.handlers, context, chat_id, name)

    def type(self, text):
        asyncio.run(self.handlers["message"](self._update(text=text), self.context))

    def screen(self):
        """The buttons on screen right now, to tap them again later (a double tap, an old menu)."""
        return self.messages[-1][1]

    def press(self, label, on=None):
        markup = on if on is not None else self.messages[-1][1]
        buttons = [b for row in markup.inline_keyboard for b in row] if markup else []
        matches = [b for b in buttons if b.text == label]
        assert matches, f"no button {label!r}; have {[b.text for b in buttons]}"
        chat = self

        class Query:
            data = matches[0].callback_data

            async def answer(self, text=None, show_alert=False):
                if text:
                    chat.alerts.append(text)

            async def edit_message_text(self, text, reply_markup=None):
                chat.messages[-1] = [text, reply_markup]

        prefix = Query.data.split(":")[0]
        if prefix == "us" and Query.data[3:4] == "t" and Query.data[4:5].isdigit():
            prefix = "us_team"
        handler = self.handlers[prefix]
        asyncio.run(handler(self._update(query=Query()), self.context))


def make_bot(tmp_path, sheets=None, lang="ru"):
    from sheetbot import access, analysis_builder, bot, builder, report_builder
    key = tmp_path / "key.json"
    key.write_text('{"client_email": "sheetbot@demo.iam.gserviceaccount.com"}')
    cfg = Config("1:x", key, "id", ZoneInfo("Asia/Bangkok"), tmp_path)
    storage = Storage(tmp_path / "db.sqlite")
    storage.claim_admin(100, "Danga", lang)
    context = SimpleNamespace(
        application=SimpleNamespace(bot_data={"config": cfg, "storage": storage,
                                              "sheets": sheets or FakeSheets()}),
        user_data={},
        bot=FakeBot(),
        job_queue=None,
        args=[],
    )
    handlers = {"message": bot.on_message, "nr": builder.on_builder_button, "rl": builder.on_rules_button,
                "rb": report_builder.on_button, "rp": report_builder.on_card_button,
                "ab": analysis_builder.on_button, "an": analysis_builder.on_card_button,
                "ac": access.on_request_button, "us": access.on_users_button,
                "us_team": access.on_user_team_button, "pz": access.on_pause_button}
    return Chat(handlers, context), storage
